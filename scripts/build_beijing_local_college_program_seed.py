"""Build first-look Beijing local college program seed from the existing workbook.

This seed is for the non-commercial V1 recommendation experience. It preserves
source and caution fields from the workbook and marks records as
local_workbook_firstlook_seed instead of pretending they are official admission
groups.
"""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT.parent / "output" / "spreadsheet" / "beijing_色弱_本科规划_2026.xlsx"
OUT_JS = ROOT / "miniprogram" / "data" / "beijing-local-college-programs.js"
OUT_REPORT = ROOT / "docs" / "beijing_local_college_program_seed_report.md"
OFFICIAL_GROUPS_JS = ROOT / "miniprogram" / "data" / "college-admission-groups.js"

PRIVATE_KEYWORDS = ("民办", "独立学院", "职业", "专科", "高职")

SUBJECT_ALIASES = {
    "政治": "思想政治",
    "思想政治": "思想政治",
    "物理": "物理",
    "化学": "化学",
    "生物": "生物",
    "历史": "历史",
    "地理": "地理",
}


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


def to_int(value):
    text = clean(value).replace(",", "")
    if not text or text in {"-", "未公开"}:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def normalize_subject_requirement(text):
    raw = clean(text)
    if not raw or "不限" in raw:
        return {
            "raw": raw or "不限选考科目",
            "subjects": [],
            "mode": "unlimited",
            "displayText": "不限",
        }
    subjects = []
    for key, label in SUBJECT_ALIASES.items():
        if key in raw and label not in subjects:
            subjects.append(label)
    mode = "all_required" if subjects else "custom"
    display = "＋".join(subjects) if subjects else raw
    return {
        "raw": raw,
        "subjects": subjects,
        "mode": mode,
        "displayText": display,
    }


def is_excluded(row):
    layer = clean(row.get("学校层次"))
    name = clean(row.get("大学名称"))
    train_type = clean(row.get("培养类型"))
    blob = " ".join([layer, name, train_type])
    return any(k in blob for k in PRIVATE_KEYWORDS)


def classify_restrictions(rows):
    color_values = {clean(r.get("色弱是否可以")) for r in rows if clean(r.get("色弱是否可以"))}
    vision_values = {clean(r.get("对视力是否有要求")) for r in rows if clean(r.get("对视力是否有要求"))}
    notes = [clean(r.get("限制说明")) for r in rows if clean(r.get("限制说明"))]
    has_limited = any("明确受限" in x or "不可以" in x for x in color_values)
    has_caution = any("谨慎" in x or "受限" in x for x in color_values)
    if has_limited:
        color_risk = "limited"
    elif has_caution:
        color_risk = "caution"
    elif color_values:
        color_risk = "allowed_or_not_found"
    else:
        color_risk = "unknown"
    return {
        "colorWeaknessRisk": color_risk,
        "hasMedicalRestriction": any("有要求" in x for x in vision_values) or any("体检" in x or "色" in x for x in notes),
        "restrictionSummary": notes[0][:120] if notes else "以官方招生目录和学校要求为准。",
    }


def major_quality(name):
    score = 0
    if "方向" in name:
        score += 2
    if name.endswith(")") or name.endswith("）"):
        score += 1
    if "类" in name:
        score -= 1
    return score


def pick_majors(rows, limit=8):
    majors = []
    seen = set()
    # Prefer allowed / no explicit conflict rows for first-look display, while
    # keeping workbook order as much as possible.
    indexed_rows = list(enumerate(rows))
    sorted_rows = sorted(indexed_rows, key=lambda pair: (
        "明确受限" in clean(pair[1].get("色弱是否可以")),
        major_quality(clean(pair[1].get("专业名称"))),
        pair[0],
    ))
    for _, row in sorted_rows:
        name = clean(row.get("专业名称"))
        if not name or name in seen:
            continue
        seen.add(name)
        majors.append({
            "majorName": name,
            "discipline": clean(row.get("学科门类")) or "待核对",
            "colorWeaknessStatus": clean(row.get("色弱是否可以")) or "以学校要求为准",
            "notes": clean(row.get("限制说明"))[:100],
        })
        if len(majors) >= limit:
            break
    return majors


def read_official_school_names():
    text = OFFICIAL_GROUPS_JS.read_text(encoding="utf-8")
    return set(re.findall(r'"collegeName":\s*"([^"]+)"', text))


def main():
    if not WORKBOOK.exists():
        raise FileNotFoundError(WORKBOOK)
    wb = load_workbook(WORKBOOK, read_only=True, data_only=True)
    ws = wb["原始数据"]
    headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    grouped = defaultdict(list)
    source_rows = 0
    excluded_rows = 0
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = dict(zip(headers, values))
        if clean(row.get("城市")) != "北京" or clean(row.get("招生省份")) != "北京":
            excluded_rows += 1
            continue
        if is_excluded(row):
            excluded_rows += 1
            continue
        college = clean(row.get("大学名称"))
        major = clean(row.get("专业名称"))
        rank = to_int(row.get("最低位次"))
        score = to_int(row.get("录取分数线"))
        if not college or not major or not rank:
            excluded_rows += 1
            continue
        source_rows += 1
        subject = normalize_subject_requirement(row.get("选科要求"))
        key = (college, subject["displayText"], score, rank)
        grouped[key].append(row)

    records = []
    for index, ((college, subject_display, score, rank), rows) in enumerate(sorted(grouped.items(), key=lambda item: (item[0][3], item[0][0], item[0][1])), start=1):
        first = rows[0]
        subject = normalize_subject_requirement(first.get("选科要求"))
        restrictions = classify_restrictions(rows)
        majors = pick_majors(rows)
        if not majors:
            continue
        record = {
            "id": f"local_{index:04d}",
            "year": 2025,
            "province": "北京",
            "batch": "本科普通批",
            "collegeName": college,
            "collegeLevel": clean(first.get("学校层次")),
            "groupName": "专业方向组/首轮筛选组",
            "subjectRequirement": subject,
            "minScore": score,
            "minRank": rank,
            "majorNames": [m["majorName"] for m in majors],
            "majors": majors,
            "sourceRowCount": len(rows),
            "source": {
                "publisher": "北京教育考试院 高校选考查询 + 北京教育考试院2025本科普通批投档线 + 北京2025高考分数分布 + 教育部体检指导意见",
                "title": "北京本地院校色弱本科志愿规划工作簿",
                "workbook": str(WORKBOOK),
                "url": clean(first.get("来源链接")),
            },
            "dataStatus": "local_workbook_firstlook_seed",
            "limitations": [
                "由本地规划工作簿聚合为首轮筛选专业方向，不伪造院校专业组代码。",
                "正式填报前需以当年官方招生目录、院校专业组和学校要求为准。",
            ],
            **restrictions,
        }
        records.append(record)

    OUT_JS.write_text("module.exports = " + json.dumps(records, ensure_ascii=False, indent=2) + ";\n", encoding="utf-8")

    official_schools = read_official_school_names()
    local_schools = {r["collegeName"] for r in records}
    added_schools = sorted(local_schools - official_schools)
    overlap_schools = sorted(local_schools & official_schools)
    lines = [
        "# 北京本地院校专业方向 seed 构建报告",
        "",
        "日期：2026-05-11",
        "",
        f"- 源工作簿：`{WORKBOOK}`",
        f"- 源专业记录数：{source_rows}",
        f"- 排除/缺字段记录数：{excluded_rows}",
        f"- 输出候选组数：{len(records)}",
        f"- 覆盖学校数：{len(local_schools)}",
        f"- 与官方投档线种子重叠学校数：{len(overlap_schools)}",
        f"- 相比当前小程序官方种子新增学校数：{len(added_schools)}",
        "",
        "## 新增学校",
        "",
    ]
    lines.extend([f"- {name}" for name in added_schools] or ["- 无"])
    lines.extend([
        "",
        "## 数据状态",
        "",
        "本 seed 标记为 `local_workbook_firstlook_seed`，用于首版非商业化初筛推荐。它不伪造院校专业组代码，展示为“专业方向组/首轮筛选组”。",
    ])
    OUT_REPORT.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"source_rows={source_rows}")
    print(f"records={len(records)}")
    print(f"schools={len(local_schools)}")
    print(f"added_schools={len(added_schools)}")
    print(f"out={OUT_JS}")
    print(f"report={OUT_REPORT}")


if __name__ == "__main__":
    main()
